#!/usr/bin/env python3
"""Frozen BIT V3 raw-capacity parser for the v0.2.1 amendment.

This source is frozen before BIT capacity access and MUST NOT be executed until
an immutable BIT registration receipt exists.

Physical-cycle mapping:
- first20 workbook local 1..20 -> physical 1..20
- first20 workbook local 21 -> excluded terminal workbook bucket
- later workbook local k -> physical 20 + k

It outputs raw discharge capacity in Ah only. It does not compute SOH, EOL,
predictions, errors, statistical tests, or model outputs.
"""

from __future__ import annotations

import argparse
import csv
from io import BytesIO
import json
import math
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook


TARGET_SHEET = "记录表"
REQUIRED_COLUMNS = (
    "Data_Point",
    "Date_Time",
    "Current(A)",
    "Capacity(Ah)",
    "Cycle_Index",
)
COMMIT_RE = re.compile(r"^[0-9a-f]{40}$", re.IGNORECASE)
SHA_RE = re.compile(r"^[0-9a-f]{64}$", re.IGNORECASE)


def read_json(path: Path) -> dict:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def as_bool(value: object) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes"}


def validate_registration_receipt(path: Path) -> dict:
    """Fail before opening the BIT archive unless immutable registration passed."""
    receipt = read_json(path)
    if receipt.get("status") != "PASS":
        raise RuntimeError("BIT immutable registration receipt is not PASS")

    commit = str(
        receipt.get("commit")
        or receipt.get("git_commit")
        or receipt.get("registered_commit")
        or ""
    )
    github_release = str(
        receipt.get("github_release")
        or receipt.get("github_release_url")
        or ""
    )
    osf_registration = str(
        receipt.get("osf_registration")
        or receipt.get("osf_registration_url")
        or ""
    )
    zenodo_doi = str(
        receipt.get("zenodo_specific_version_doi")
        or receipt.get("zenodo_doi")
        or receipt.get("doi")
        or ""
    )
    freeze_sha = str(
        receipt.get("freeze_zip_sha256")
        or receipt.get("freeze_sha256")
        or receipt.get("bit_freeze_sha256")
        or ""
    )

    if not COMMIT_RE.fullmatch(commit):
        raise RuntimeError("Registration receipt lacks a 40-character commit")
    if not github_release.startswith("https://github.com/"):
        raise RuntimeError("Registration receipt lacks a GitHub release URL")
    if "osf.io/" not in osf_registration:
        raise RuntimeError("Registration receipt lacks an OSF registration URL")
    if not zenodo_doi.startswith("10.5281/zenodo."):
        raise RuntimeError("Registration receipt lacks a Zenodo version DOI")
    if not SHA_RE.fullmatch(freeze_sha):
        raise RuntimeError("Registration receipt lacks the frozen ZIP SHA-256")
    if receipt.get("bit_evaluation_unlocked") is not True:
        raise RuntimeError("Registration receipt has not unlocked BIT evaluation")
    return receipt


def finite_float(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def normalize_header(value: object) -> str:
    return "" if value is None else str(value).strip()


def discharge_capacity_for_cycle(rows: list[tuple[float, float]]) -> float | None:
    """Return the span of cumulative Capacity(Ah) during negative-current rows."""
    capacities = [
        capacity
        for current, capacity in rows
        if current < -1e-9 and math.isfinite(capacity)
    ]
    if len(capacities) < 2:
        return None
    span = max(capacities) - min(capacities)
    if not math.isfinite(span) or span <= 0:
        return None
    return float(span)


def extract_workbook_member(
    archive: zipfile.ZipFile,
    member: str,
    role: str,
) -> list[dict[str, object]]:
    workbook_bytes = archive.read(member)
    workbook = load_workbook(
        BytesIO(workbook_bytes),
        read_only=True,
        data_only=True,
    )
    try:
        if TARGET_SHEET not in workbook.sheetnames:
            raise ValueError(f"Worksheet {TARGET_SHEET!r} absent: {member}")
        sheet = workbook[TARGET_SHEET]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError(f"Empty worksheet: {member}")

        header_map = {
            normalize_header(value): index
            for index, value in enumerate(header)
        }
        missing = [
            column for column in REQUIRED_COLUMNS
            if column not in header_map
        ]
        if missing:
            raise ValueError(f"Missing columns {missing}: {member}")

        grouped: dict[int, list[tuple[float, float]]] = {}
        for row in rows:
            cycle_value = finite_float(row[header_map["Cycle_Index"]])
            current = finite_float(row[header_map["Current(A)"]])
            capacity = finite_float(row[header_map["Capacity(Ah)"]])
            if cycle_value is None or current is None or capacity is None:
                continue
            local_cycle = int(round(cycle_value))
            if role == "first20":
                if not 1 <= local_cycle <= 20:
                    continue
                physical_cycle = local_cycle
            elif role == "later":
                if local_cycle < 1:
                    continue
                physical_cycle = 20 + local_cycle
            else:
                raise ValueError(f"Unknown workbook role: {role}")
            grouped.setdefault(physical_cycle, []).append((current, capacity))

        extracted: list[dict[str, object]] = []
        for physical_cycle in sorted(grouped):
            capacity_ah = discharge_capacity_for_cycle(grouped[physical_cycle])
            extracted.append(
                {
                    "physical_cycle": physical_cycle,
                    "raw_discharge_capacity_Ah": capacity_ah,
                    "source_member": member,
                    "source_role": role,
                }
            )
        return extracted
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--archive", type=Path, required=True)
    parser.add_argument("--pairing-manifest", type=Path, required=True)
    parser.add_argument("--mapping-manifest", type=Path, required=True)
    parser.add_argument("--registration-receipt", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    # Deliberately before every archive open.
    validate_registration_receipt(args.registration_receipt)

    pairing_rows = read_csv(args.pairing_manifest)
    mapping_rows = read_csv(args.mapping_manifest)
    mapping_by_cell = {row["cell_id"]: row for row in mapping_rows}

    outputs: list[dict[str, object]] = []
    with zipfile.ZipFile(args.archive) as archive:
        for pair in sorted(pairing_rows, key=lambda row: int(row["cell_number"])):
            cell_id = pair["cell_id"]
            mapping = mapping_by_cell.get(cell_id)
            if mapping is None:
                raise ValueError(f"Missing mapping row: {cell_id}")
            if not as_bool(mapping["include_in_amended_BIT_evaluation"]):
                continue

            first20_member = pair["first20_member"]
            later_member = pair["later_member"]
            if not first20_member or not later_member:
                raise ValueError(f"Included cell lacks paired workbooks: {cell_id}")

            cell_records = []
            cell_records.extend(
                extract_workbook_member(archive, first20_member, "first20")
            )
            cell_records.extend(
                extract_workbook_member(archive, later_member, "later")
            )

            seen: set[int] = set()
            for record in sorted(
                cell_records,
                key=lambda item: int(item["physical_cycle"]),
            ):
                physical_cycle = int(record["physical_cycle"])
                if physical_cycle in seen:
                    raise RuntimeError(
                        f"Duplicate physical cycle {physical_cycle}: {cell_id}"
                    )
                seen.add(physical_cycle)
                capacity = record["raw_discharge_capacity_Ah"]
                quality_pass = bool(
                    capacity is not None and 0 < float(capacity) <= 3.6
                )
                outputs.append(
                    {
                        "cell_id": cell_id,
                        "cell_number": pair["cell_number"],
                        "cohort": pair["cohort"],
                        "physical_cycle": physical_cycle,
                        "raw_discharge_capacity_Ah": (
                            "" if capacity is None else capacity
                        ),
                        "capacity_quality_pass": quality_pass,
                        "source_member": record["source_member"],
                        "source_role": record["source_role"],
                    }
                )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "cell_id",
        "cell_number",
        "cohort",
        "physical_cycle",
        "raw_discharge_capacity_Ah",
        "capacity_quality_pass",
        "source_member",
        "source_role",
    ]
    with args.output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(outputs)

    print(f"[PASS] raw BIT capacity rows={len(outputs)}")
    print(f"OUTPUT={args.output}")


if __name__ == "__main__":
    main()
